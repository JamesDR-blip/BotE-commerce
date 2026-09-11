"""Bounded-memory streaming readers for supported tabular sources."""

from __future__ import annotations

from io import BytesIO
from pathlib import PurePath
from typing import BinaryIO, Iterator

import pandas as pd
from openpyxl import load_workbook

from infrastructure.upload.tabular_reader import UploadValidationError


class StreamingTabularReader:
    """Read CSV/XLSX sources in bounded DataFrame chunks.

    XLS remains on the legacy reader path because openpyxl does not support
    the legacy XLS format.
    """

    _SUPPORTED_EXTENSIONS = {".csv", ".xlsx"}
    _FORMULA_PREFIXES = ("=", "+", "-", "@")

    def __init__(self, max_bytes: int, chunk_size: int = 10_000) -> None:
        if max_bytes <= 0:
            raise ValueError("max_bytes must be positive")
        if chunk_size <= 0:
            raise ValueError("chunk_size must be positive")

        self._max_bytes = max_bytes
        self._chunk_size = chunk_size

    def stream(
        self,
        filename: str | None,
        source: bytes | BinaryIO,
    ) -> Iterator[pd.DataFrame]:
        safe_name = PurePath(filename or "").name
        extension = PurePath(safe_name).suffix.lower()

        if (
            not safe_name
            or safe_name != filename
            or extension not in self._SUPPORTED_EXTENSIONS
        ):
            raise UploadValidationError(
                "Only CSV and XLSX uploads support streaming"
            )

        if isinstance(source, bytes):
            if not source:
                raise UploadValidationError("Input file is empty")
            if len(source) > self._max_bytes:
                raise UploadValidationError(
                    f"Input file exceeds the {self._max_bytes}-byte limit"
                )
            source = BytesIO(source)
        else:
            self._validate_binary_source(source)

        if extension == ".csv":
            yield from self._stream_csv(source)
        else:
            yield from self._stream_xlsx(source)

    def _validate_binary_source(self, source: BinaryIO) -> None:
        """Enforce the upload size ceiling for file-backed sources."""
        try:
            current_position = source.tell()
            source.seek(0, 2)
            size = source.tell()
            source.seek(current_position)
        except (AttributeError, OSError) as exc:
            raise UploadValidationError(
                "Streaming source must be a seekable binary file"
            ) from exc

        if size <= 0:
            raise UploadValidationError("Input file is empty")

        if size > self._max_bytes:
            raise UploadValidationError(
                f"Input file exceeds the {self._max_bytes}-byte limit"
            )

    def _stream_csv(self, source: BinaryIO) -> Iterator[pd.DataFrame]:
        try:
            reader = pd.read_csv(source, chunksize=self._chunk_size)
            yielded = False

            for chunk in reader:
                if chunk.empty:
                    continue

                yielded = True
                yield self._escape_spreadsheet_formulas(chunk)

            if not yielded:
                raise UploadValidationError("Input has no records")
        except UploadValidationError:
            raise
        except Exception as exc:
            raise UploadValidationError(
                "Could not parse the uploaded CSV file"
            ) from exc

    def _stream_xlsx(self, source: BinaryIO) -> Iterator[pd.DataFrame]:
        workbook = None

        try:
            workbook = load_workbook(
                source,
                read_only=True,
                data_only=False,
            )
            worksheet = workbook.active
            rows = worksheet.iter_rows(values_only=True)

            try:
                headers = next(rows)
            except StopIteration as exc:
                raise UploadValidationError(
                    "Input has no records"
                ) from exc

            headers = list(headers)
            if not headers or all(value is None for value in headers):
                raise UploadValidationError("Input has no records")

            batch: list[tuple[object, ...]] = []

            for row in rows:
                batch.append(tuple(row))

                if len(batch) >= self._chunk_size:
                    frame = pd.DataFrame(batch, columns=headers)
                    if not frame.empty:
                        yield self._escape_spreadsheet_formulas(frame)
                    batch = []

            if batch:
                frame = pd.DataFrame(batch, columns=headers)
                if not frame.empty:
                    yield self._escape_spreadsheet_formulas(frame)

        except UploadValidationError:
            raise
        except Exception as exc:
            raise UploadValidationError(
                "Could not parse the uploaded XLSX file"
            ) from exc
        finally:
            if workbook is not None:
                workbook.close()

    def _escape_spreadsheet_formulas(
        self,
        frame: pd.DataFrame,
    ) -> pd.DataFrame:
        sanitized = frame.copy()

        for column in sanitized.select_dtypes(
            include=["object", "string"]
        ).columns:
            sanitized[column] = sanitized[column].map(
                self._escape_value
            )

        return sanitized

    def _escape_value(self, value: object) -> object:
        if isinstance(value, str) and value.startswith(
            self._FORMULA_PREFIXES
        ):
            return "'" + value
        return value


__all__ = ["StreamingTabularReader"]
